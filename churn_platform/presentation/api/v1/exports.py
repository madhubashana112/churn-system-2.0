"""Download the latest stored analysis, with the same filters as the customer table."""
import csv
import io
import json
import base64
import zipfile
import re
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from churn_platform.presentation.api.auth import require_tenant
from churn_platform.presentation.api.dependencies import get_analysis_repo, get_tenant_repo

router = APIRouter(prefix="/exports", tags=["Exports"], dependencies=[Depends(require_tenant)])


def spreadsheet_text(value):
    """Treat uploaded values as text, never spreadsheet formulas."""
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        text = "'" + text
    return ILLEGAL_CHARACTERS_RE.sub("", text)


@router.get("")
async def export_analysis(
    tenant_id: str,
    format: Literal["csv", "xlsx", "pdf", "zip"] = "csv",
    tier: Literal["ALL", "CRITICAL", "HIGH", "MEDIUM", "LOW"] = "ALL",
    search: str = Query("", max_length=200),
):
    run = await get_analysis_repo().latest(tenant_id)
    if run is None:
        raise HTTPException(404, "Run an analysis before exporting data")
    if format == "zip":
        if not run.original_files:
            raise HTTPException(404, "Original files are unavailable for this older analysis. Upload and analyze the files again to enable this download.")
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index, original in enumerate(run.original_files, 1):
                name = original.filename.replace("\\", "/").rsplit("/", 1)[-1]
                name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(". ") or "upload"
                archive.writestr(f"{index:03d}/{name}", base64.b64decode(original.content_base64))
        return Response(buffer.getvalue(), media_type="application/zip", headers={
            "Content-Disposition": 'attachment; filename="original-uploads.zip"', "Cache-Control": "no-store"})
    selected = [o for o in sorted(run.outcomes, key=lambda o: o.prediction.churn_probability, reverse=True)
                if (tier == "ALL" or o.prediction.risk_tier == tier) and search.lower() in o.prediction.entity_id.lower()]
    if format == "pdf":
        from churn_platform.presentation.pdf_report import analysis_pdf
        from starlette.concurrency import run_in_threadpool
        tenant = await get_tenant_repo().get(tenant_id)
        content = await run_in_threadpool(analysis_pdf, run, selected, tenant.name if tenant else run.sector, tier, search)
        return Response(content, media_type="application/pdf", headers={
            "Content-Disposition": 'attachment; filename="churn-analysis.pdf"', "Cache-Control": "no-store"})
    rows = [["Customer ID", "Churn probability", "Risk tier", "Primary drivers", "Root cause",
             "Action", "Channel", "Recommendation", "Sector", "Analyzed at (UTC)", "Engine", "Evidence (JSON)"]]
    for outcome in selected:
        p, action = outcome.prediction, outcome.playbook
        if (tier != "ALL" and p.risk_tier != tier) or search.lower() not in p.entity_id.lower():
            continue
        row = [p.entity_id, p.churn_probability, p.risk_tier, "; ".join(p.primary_drivers or []), p.root_cause,
               action.action_type, action.channel, action.action_payload, run.sector, run.created_at.isoformat(),
               "System (local)" if run.offline_mode else "AI", json.dumps(outcome.features, ensure_ascii=False)]
        rows.append([v if isinstance(v, (float, int)) else spreadsheet_text(v) for v in row])
    if format == "csv":
        output = io.StringIO(newline="")
        csv.writer(output).writerows(rows)
        content = output.getvalue().encode("utf-8-sig")
        media = "text/csv; charset=utf-8"
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Customer risk"
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        from openpyxl.styles import Font, PatternFill
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F46E5")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(20, len(str(column[0].value))+3))
        for cell in sheet["B"][1:]:
            cell.number_format = "0.0%"
        output = io.BytesIO()
        book.save(output)
        content = output.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    filename = f"churn-analysis-{run.created_at:%Y%m%d}-{tier.lower()}.{format}"
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})
