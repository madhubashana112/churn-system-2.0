import asyncio
import csv
import io
import time
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from churn_platform.main import app
from churn_platform.presentation.api.auth import COOKIE, digest
from churn_platform.infrastructure.repositories.state_store import store, StateStore


@pytest.fixture
def isolated_store(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "path", tmp_path / "test.sqlite3")
    return store


def signup(client, email=None):
    payload = {"name": "Alex", "email": email or f"{uuid4().hex}@example.com", "password": "A-long-test-password"}
    response = client.post("/api/auth/signup", json=payload)
    assert response.status_code == 201, response.text
    return payload, response


def workspace(client):
    response = client.post("/api/v1/tenants/", json={"name": "Acme", "sector": "SaaS"})
    assert response.status_code == 200
    return response.json()["tenant_id"]


def test_signup_login_logout_and_persistence(isolated_store):
    client = TestClient(app)
    assert client.get("/dashboard", follow_redirects=False).headers["location"] == "/login"
    assert client.get("/api/v1/analytics/status").status_code == 401
    payload, response = signup(client, "Alex@Example.com")
    assert "httponly" in response.headers["set-cookie"].lower()
    assert "samesite=strict" in response.headers["set-cookie"].lower()
    tenant = workspace(client)
    token = client.cookies[COOKIE]
    assert client.post("/api/auth/logout").status_code == 200
    client.cookies.set(COOKIE, token)
    assert client.get("/api/auth/me").status_code == 401
    client.cookies.clear()
    assert client.post("/api/auth/login", json={**payload, "password": "incorrect"}).status_code == 401
    assert client.post("/api/auth/login", json=payload).status_code == 200
    assert client.get("/api/auth/me").json()["email"] == "alex@example.com"
    assert tenant in client.get("/dashboard", follow_redirects=False).headers["location"]
    assert client.post("/api/auth/signup", json=payload).status_code == 409
    fresh = StateStore()
    fresh.path = store.path
    saved = asyncio.run(fresh.get("account:" + digest("alex@example.com")))
    assert saved and payload["password"] not in str(saved)
    assert asyncio.run(fresh.get("tenant:" + tenant))["name"] == "Acme"


def test_validation_expiry_csrf_and_throttling(isolated_store):
    client = TestClient(app)
    assert client.post("/api/auth/signup", json={"name":" ", "email":"bad", "password":"tiny"}).status_code == 422
    payload, _ = signup(client)
    assert client.post("/api/auth/logout", headers={"Origin":"https://attacker.example"}).status_code == 403
    assert client.post("/api/auth/logout", headers={"Sec-Fetch-Site":"cross-site"}).status_code == 403
    assert client.get("/api/auth/me").status_code == 200
    with store.connect() as db:
        db.execute("UPDATE state SET expires=? WHERE key LIKE 'session:%'", (time.time()-1,))
    assert client.get("/api/auth/me").status_code == 401
    for _ in range(14):
        assert client.post("/api/auth/login", json={**payload,"password":"wrong"}).status_code == 401
    assert client.post("/api/auth/login", json=payload).status_code == 429


def test_other_account_cannot_read_write_or_export(isolated_store):
    owner, stranger = TestClient(app), TestClient(app)
    signup(owner)
    tenant = workspace(owner)
    signup(stranger)
    for url in [f"/api/v1/tenants/{tenant}", f"/dashboard?tenant_id={tenant}",
                f"/customer?tenant_id={tenant}&entity_id=x", f"/api/v1/analytics/metrics?tenant_id={tenant}",
                f"/api/v1/analytics/customer?tenant_id={tenant}&entity_id=x",
                f"/api/v1/upload/demo-data?tenant_id={tenant}", f"/api/v1/exports?tenant_id={tenant}"]:
        assert stranger.get(url).status_code == 404, url
    assert stranger.post("/api/v1/upload/analyze", data={"tenant_id":tenant}, files={"files":("users.csv",b"user_id\nx")}).status_code == 404
    assert TestClient(app).get(f"/api/v1/exports?tenant_id={tenant}").status_code == 401


def test_exports_include_all_rows_filters_and_safe_spreadsheet_values(isolated_store):
    from churn_platform.domain.models.analysis_run import AnalysisRun, EntityOutcome
    from churn_platform.domain.models.churn_prediction import ChurnPrediction
    from churn_platform.domain.models.retention_playbook import RetentionPlaybook
    from churn_platform.presentation.api.dependencies import get_analysis_repo
    from conftest import sector_schema
    client = TestClient(app)
    signup(client)
    tenant = workspace(client)
    assert client.get(f"/api/v1/exports?tenant_id={tenant}").status_code == 404
    run = AnalysisRun(tenant_id=tenant, sector="SaaS", schema_mapping=sector_schema("SaaS"), offline_mode=True,
        outcomes=[EntityOutcome(prediction=ChurnPrediction(entity_id=f"user_{i}", churn_probability=.9 if i==0 else .2,
            risk_tier="HIGH" if i==0 else "LOW", root_cause='=HYPERLINK("evil")'),
            playbook=RetentionPlaybook(action_type="EMAIL", channel="Email", action_payload="+formula"), features={"usage":i}) for i in range(300)])
    asyncio.run(get_analysis_repo().save(run))
    response = client.get(f"/api/v1/exports?tenant_id={tenant}")
    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 301
    assert rows[1][4].startswith("'=") and rows[1][7] == "'+formula"
    assert "attachment" in response.headers["content-disposition"]
    filtered = client.get(f"/api/v1/exports?tenant_id={tenant}&format=xlsx&tier=HIGH&search=user_0")
    book = load_workbook(io.BytesIO(filtered.content))
    assert book.active.max_row == 2
    assert book.active["B2"].value == .9
    assert book.active["E2"].data_type == "s"
    from churn_platform.presentation.api.v1.exports import spreadsheet_text
    assert spreadsheet_text("=bad\x00") == "'=bad"
    assert len(list(csv.reader(io.StringIO(client.get(f"/api/v1/exports?tenant_id={tenant}&search=absent").text)))) == 1
    assert client.get(f"/api/v1/exports?tenant_id={tenant}&format=exe").status_code == 422


def test_auth_pages_and_shared_controls(isolated_store):
    client = TestClient(app)
    assert 'autocomplete="current-password"' in client.get("/login").text
    assert 'autocomplete="new-password"' in client.get("/signup").text
    signup(client)
    tenant = workspace(client)
    page = client.get(f"/dashboard?tenant_id={tenant}")
    assert page.status_code == 200
    assert 'data-theme-toggle' in page.text and 'data-logout' in page.text
    assert 'data-export="csv"' in page.text and 'data-export="xlsx"' in page.text
    assert page.headers["cache-control"] == "no-store"


def test_login_always_lands_on_tenant_registration_with_existing_workspace(isolated_store):
    client = TestClient(app)
    credentials, created = signup(client)
    assert created.json()["redirect"] == "/"
    tenant = workspace(client)
    client.post("/api/auth/logout")
    logged_in = client.post("/api/auth/login", json=credentials)
    assert logged_in.status_code == 200
    assert logged_in.json()["redirect"] == "/"
    landing = client.get(logged_in.json()["redirect"], follow_redirects=False)
    assert landing.status_code == 200
    assert "Register your tenant" in landing.text
    for path in ("/login", "/signup"):
        assert client.get(path, follow_redirects=False).headers["location"] == "/"
    # Existing workspaces remain accessible through an explicit navigation.
    assert client.get(f"/dashboard?tenant_id={tenant}").status_code == 200
